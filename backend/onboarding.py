"""Resume review and capability URLs, without starting a brain on the home page."""
import io
import json
import secrets
import threading
import time
import zipfile
from pathlib import Path

import httpx
from fastapi import APIRouter, HTTPException, Request, UploadFile
from fastapi.responses import JSONResponse
from pydantic import BaseModel

from .jobs import fetch_resume_jobs
from .storage import STATE, Store

router = APIRouter(prefix="/api")
HEADERS = {"User-Agent": "Mozilla/5.0 Jobfly/1.0"}
conversion_slots = threading.BoundedSemaphore(2)
limits = {}
limit_lock = threading.Lock()


def throttle(request):
    # Bound expensive public conversion/matching calls; never retain documents here.
    key = request.client.host if request.client else "unknown"
    now = time.monotonic()
    with limit_lock:
        for old in list(limits):
            limits[old] = [t for t in limits[old] if now - t < 600]
            if not limits[old]:
                del limits[old]
        recent = limits.setdefault(key, [])
        if len(recent) >= 12:
            raise HTTPException(429, "A few too many requests. Try again in ten minutes.")
        recent.append(now)


def validate_resume(data):
    if not isinstance(data, dict) or not isinstance(data.get("basics"), dict):
        raise ValueError('Your JSON Resume needs a "basics" object.')
    if not isinstance(data["basics"].get("name"), str) or not data["basics"]["name"].strip():
        raise ValueError("Add your name to basics.name before continuing.")
    if data["basics"].get("label") is not None and not isinstance(data["basics"]["label"], str):
        raise ValueError("basics.label must be text.")
    if len(json.dumps(data)) > 500_000:
        raise ValueError("Keep your resume under 500 KB.")
    return data


@router.get("/thomas")
def thomas():
    try:
        with httpx.Client(timeout=30, headers=HEADERS, follow_redirects=True) as client:
            response = client.get("https://registry.jsonresume.org/thomasdavis.json")
            response.raise_for_status()
            return validate_resume(response.json())
    except Exception as exc:
        raise HTTPException(502, "Thomas’s resume could not load. Try again or upload yours.") from exc


def extract_document(content, suffix):
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        if len(reader.pages) > 30:
            raise ValueError("Choose a resume with 30 pages or fewer.")
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif suffix == ".docx":
        from docx import Document
        archive = zipfile.ZipFile(io.BytesIO(content))
        if sum(info.file_size for info in archive.infolist()) > 30_000_000:
            raise ValueError("This Word document is too large to unpack.")
        doc = Document(io.BytesIO(content))
        text = "\n".join([p.text for p in doc.paragraphs] + [cell.text for table in doc.tables for row in table.rows for cell in row.cells])
    elif suffix == ".txt":
        text = content.decode("utf-8-sig")
    else:
        raise ValueError("Upload a PDF, DOCX, TXT, or JSON file.")
    if len(text.strip()) < 40:
        raise ValueError("We couldn’t read enough text. For a scanned PDF, upload a text version instead.")
    if len(text) > 60_000:
        raise ValueError("This document is too long. Upload just your resume.")
    return text


@router.post("/convert")
def convert(request: Request, file: UploadFile):
    throttle(request)
    content = file.file.read(5_000_001)
    if len(content) > 5_000_000:
        raise HTTPException(413, "Choose a file under 5 MB.")
    try:
        suffix = Path(file.filename or "").suffix.lower()
        if suffix == ".json":
            return validate_resume(json.loads(content))
        text = extract_document(content, suffix)
        if not conversion_slots.acquire(blocking=False):
            raise HTTPException(429, "Two resumes are being converted. Try again in a moment.")
        try:
            from .resume_conversion import convert_text
            return validate_resume(convert_text(text))
        finally:
            conversion_slots.release()
    except HTTPException:
        raise
    except (ValueError, UnicodeError, zipfile.BadZipFile) as exc:
        raise HTTPException(400, str(exc)) from exc
    except Exception as exc:
        raise HTTPException(502, "Resume conversion is temporarily unavailable. Try again, or upload JSON Resume.") from exc


class StartSession(BaseModel):
    resume: dict


@router.post("/sessions")
def start(body: StartSession, request: Request):
    throttle(request)
    try:
        data = validate_resume(body.resume)
        jobs = fetch_resume_jobs(data)
        identifier = secrets.token_hex(24)
        store = Store(STATE / identifier)
        try:
            store.save(dict(resume=data, jobs=jobs, source=jobs[0]["source"], marks={}, reasons={}))
        finally:
            store.db.close()
        return {"url": f"/s/{identifier}", "count": len(jobs)}
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    except Exception as exc:
        raise HTTPException(502, "Job search is temporarily unavailable. Your resume is still here; try again.") from exc


def session_exists(identifier):
    return (len(identifier) == 48 and all(c in "0123456789abcdef" for c in identifier)
            and (STATE / identifier / "jobfly.sqlite").is_file())


@router.get("/resume.json")
def download(request: Request):
    store = Store(STATE / request.state.session_id)
    try:
        data = store.latest()
        if not data or not data.get("resume"):
            raise HTTPException(404, "This session has no resume yet.")
        return JSONResponse(data["resume"], headers={"Content-Disposition": 'attachment; filename="resume.json"'})
    finally:
        store.db.close()
