import io
import json
from types import SimpleNamespace

import pytest
from docx import Document
from fastapi.testclient import TestClient

from backend import onboarding, server
from backend.jobs import SensoryEncoder
from backend.storage import Store


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(onboarding, "STATE", tmp_path)
    monkeypatch.setattr(server, "pool", SimpleNamespace())
    onboarding.limits.clear()
    return TestClient(server.app)


def test_session_url_restores_resume_and_has_no_examples(client, monkeypatch, tmp_path):
    jobs = [{"id": "real-1", "title": "Engineer", "company": "Employer", "source": "arbeitnow", "url": "https://www.arbeitnow.com/jobs/example"}]
    monkeypatch.setattr(onboarding, "fetch_resume_jobs", lambda _: jobs)
    resume = {"basics": {"name": "Test Person"}}
    first = client.post("/api/sessions", json={"resume": resume}).json()
    second = client.post("/api/sessions", json={"resume": resume}).json()
    assert first["url"] != second["url"]
    identifier = first["url"].split("/")[-1]
    store = Store(tmp_path / identifier)
    assert store.latest()["jobs"] == jobs
    store.db.close()
    response = client.get("/api/resume.json", params={"session": identifier})
    assert response.json() == resume
    assert response.headers["cache-control"] == "no-store"
    assert client.get("/api/resume.json", params={"session": "a" * 48}).status_code == 404
    assert client.get("/api/state").status_code == 401
    assert client.post("/api/examples", params={"session": identifier}).status_code in (404, 405)
    assert client.post("/api/sessions", json={"resume": resume}, headers={"Origin": "https://other.example"}).status_code == 403


def test_invalid_resume_creates_no_session(client, tmp_path):
    response = client.post("/api/sessions", json={"resume": {"basics": {}}})
    assert response.status_code == 400
    assert list(tmp_path.iterdir()) == []


def test_json_upload_is_lossless_and_never_calls_model(client, monkeypatch):
    from backend import resume_conversion
    monkeypatch.setattr(resume_conversion, "convert_text", lambda _: pytest.fail("JSON must not invoke an LLM"))
    resume = {"basics": {"name": "Test Person"}, "custom": {"preserved": True}}
    response = client.post("/api/convert", files={"file": ("resume.json", json.dumps(resume), "application/json")})
    assert response.json() == resume
    assert client.post("/api/convert", files={"file": ("resume.json", "bad json")}).status_code == 400


def test_word_extraction_includes_tables():
    doc = Document()
    doc.add_paragraph("Test Person — Software Engineer. Experience building accessible interfaces.")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "React and TypeScript"
    buffer = io.BytesIO()
    doc.save(buffer)
    text = onboarding.extract_document(buffer.getvalue(), ".docx")
    assert "React and TypeScript" in text
    with pytest.raises(ValueError):
        onboarding.extract_document(b"scan", ".exe")


def test_empty_habitat_has_no_synthetic_neural_input():
    encoder = SensoryEncoder([], 12)
    assert encoder.codes.shape == (0, 12)


def test_conversion_failover_keeps_zero_price_and_validated_output(monkeypatch):
    from backend import resume_conversion as conversion
    attempts = []
    parsed = conversion.Resume(basics=conversion.Basics(name="Test Person"))

    async def parse(**kwargs):
        attempts.append(kwargs)
        if len(attempts) == 1:
            raise TimeoutError()
        call = SimpleNamespace(function=SimpleNamespace(name="submit_resume", parsed_arguments=parsed))
        return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(tool_calls=[call]))])

    class Client:
        def __init__(self, **kwargs):
            self.chat = SimpleNamespace(completions=SimpleNamespace(parse=parse))

        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            pass

    monkeypatch.setattr(conversion, "AsyncOpenAI", Client)
    monkeypatch.setenv("OPENROUTER_API_KEY", "unit-test-only")
    monkeypatch.delenv("JOBFLY_CONVERSION_KEY_FILE", raising=False)
    assert conversion.convert_text("Resume text")["basics"]["name"] == "Test Person"
    assert len(attempts) == 2
    for attempt in attempts:
        assert attempt["extra_body"]["provider"]["max_price"] == {"prompt": 0, "completion": 0}
        assert attempt["tool_choice"]["function"]["name"] == "submit_resume"
